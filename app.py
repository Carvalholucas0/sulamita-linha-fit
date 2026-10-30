import io
import json
import secrets
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
from functools import wraps

from flask import Flask, Response, jsonify, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from sqlalchemy.exc import IntegrityError, SQLAlchemyError

from config import Config
from leads import normalize_whatsapp
from migrations import upgrade_legacy_schema
from models import Lead, MarketingConsentEvent, Registration, SheetSyncJob, db
from quiz import (
    ALL_PROFILES,
    CAMPAIGN_SOURCE,
    MARKETING_CONSENT_VERSION,
    QUESTIONS,
    PRIVACY_NOTICE_VERSION,
    PROFILES,
    QUIZ_VERSION,
    build_answer_snapshot,
    build_declared_preferences,
    build_score_snapshot,
    calculate_profile,
    format_preference,
    option_label,
)
from sheets import SheetSyncError, is_configured, sync_registration


def create_app(test_config=None):
    app = Flask(__name__)
    app.config.from_object(Config)
    if test_config:
        app.config.update(test_config)

    db.init_app(app)
    sheet_sync_executor = ThreadPoolExecutor(max_workers=1, thread_name_prefix="sheets-sync")
    with app.app_context():
        db.create_all()
        upgrade_legacy_schema(db.engine, app.logger)
        registrations_without_job = (
            db.session.query(Registration.id)
            .outerjoin(SheetSyncJob)
            .filter(SheetSyncJob.registration_id.is_(None))
            .all()
        )
        if registrations_without_job:
            db.session.add_all(
                SheetSyncJob(registration_id=registration_id)
                for (registration_id,) in registrations_without_job
            )
            db.session.commit()

    def retry_sheet_sync(limit):
        if not is_configured(app.config):
            return set()

        jobs = (
            SheetSyncJob.query.filter_by(synced_at=None)
            .order_by(SheetSyncJob.registration_id.asc())
            .limit(limit)
            .all()
        )
        if not jobs:
            return set()

        synced_ids = set()
        for job in jobs:
            try:
                sync_registration(job.registration, ALL_PROFILES, QUESTIONS, app.config)
            except SheetSyncError as error:
                job.attempt_count += 1
                job.last_error = str(error)[:500]
                app.logger.error(
                    "Google Sheets sync failed for registration %s (attempt %s): %s",
                    job.registration_id,
                    job.attempt_count,
                    job.last_error,
                )
            else:
                job.attempt_count += 1
                job.last_error = None
                job.synced_at = datetime.now(timezone.utc)
                synced_ids.add(job.registration_id)

        try:
            db.session.commit()
        except SQLAlchemyError:
            db.session.rollback()
            app.logger.exception("Falha ao salvar o estado de sincronização da planilha.")
            return set()
        return synced_ids

    def dispatch_sheet_sync(limit):
        if not is_configured(app.config):
            return set()
        if app.testing:
            return retry_sheet_sync(limit)

        def sync_in_app_context():
            with app.app_context():
                retry_sheet_sync(limit)

        sheet_sync_executor.submit(sync_in_app_context)
        return set()

    def validated_answers(data):
        answers = data.get("answers")
        question_ids = {question["id"] for question in QUESTIONS}
        if not isinstance(answers, dict) or set(answers) != question_ids:
            return None
        for question in QUESTIONS:
            valid_options = {option["id"] for option in question["options"]}
            if (
                not isinstance(answers[question["id"]], str)
                or answers[question["id"]] not in valid_options
            ):
                return None
        return answers

    def registration_response(registration):
        synced_ids = dispatch_sheet_sync(limit=1)
        response = jsonify(
            success=True,
            profile={"id": registration.profile, **PROFILES[registration.profile]},
            sheet_backup=(
                "synced" if registration.id in synced_ids else
                "pending" if is_configured(app.config) else
                "not_configured"
            ),
        )
        response.headers["Cache-Control"] = "no-store"
        return response

    def require_admin(view):
        @wraps(view)
        def wrapped(*args, **kwargs):
            password = app.config.get("ADMIN_PASSWORD")
            if not password:
                return "Área administrativa não configurada.", 503

            auth = request.authorization
            username = auth.username if auth else ""
            supplied_password = (auth.password or "") if auth else ""
            valid_user = secrets.compare_digest(username, app.config["ADMIN_USER"])
            valid_password = secrets.compare_digest(supplied_password, password)
            if not (valid_user and valid_password):
                return Response(
                    "Autenticação necessária.",
                    401,
                    {"WWW-Authenticate": 'Basic realm="Sulamita Admin"'},
                )
            return view(*args, **kwargs)

        return wrapped

    @app.get("/")
    def index():
        return render_template("index.html")

    @app.get("/privacidade")
    def privacy_notice():
        return render_template("privacy.html")

    @app.get("/api/questions")
    def get_questions():
        return jsonify(QUESTIONS)

    @app.post("/api/submit")
    def submit_quiz():
        data = request.get_json(silent=True)
        if not isinstance(data, dict):
            return jsonify(error="Envie os dados do cadastro em JSON."), 400

        name = data.get("name")
        whatsapp = data.get("whatsapp")
        neighborhood = data.get("neighborhood")
        consent = data.get("data_consent")
        marketing_consent = data.get("marketing_consent", False)
        submission_key = data.get("submission_key")

        if not isinstance(name, str) or not name.strip() or len(name.strip()) > 100:
            return jsonify(error="Informe seu primeiro nome.", field="name"), 400
        first_name = name.strip().split()[0]
        if not isinstance(whatsapp, str) or len(whatsapp.strip()) > 30:
            return jsonify(error="Informe um WhatsApp com DDD.", field="whatsapp"), 400
        normalized_whatsapp = normalize_whatsapp(whatsapp)
        if not normalized_whatsapp:
            return jsonify(error="Confira o número e o DDD do WhatsApp.", field="whatsapp"), 400
        if (
            not isinstance(neighborhood, str)
            or not neighborhood.strip()
            or len(neighborhood.strip()) > 160
        ):
            return jsonify(error="Informe seu bairro e cidade.", field="neighborhood"), 400
        if consent is not True:
            return jsonify(
                error="Leia e aceite o aviso para concluir seu cadastro.",
                field="data-consent",
            ), 400
        if not isinstance(marketing_consent, bool):
            return jsonify(
                error="A autorização de novidades é inválida.",
                field="marketing_consent",
            ), 400
        if not isinstance(submission_key, str) or len(submission_key) != 36:
            return jsonify(
                error="Atualize a página e tente novamente.",
                field="submission",
            ), 400
        answers = validated_answers(data)
        if answers is None:
            return jsonify(error="Responda todas as perguntas do quiz."), 400

        existing_registration = Registration.query.filter_by(
            submission_key=submission_key
        ).one_or_none()
        if existing_registration:
            return registration_response(existing_registration)

        profile_id = calculate_profile(answers)
        now = datetime.now(timezone.utc)

        def save_registration(lead=None):
            current_lead = lead or Lead.query.filter_by(
                whatsapp_normalized=normalized_whatsapp
            ).one_or_none()
            if current_lead is None:
                current_lead = Lead(
                    name=first_name,
                    whatsapp=whatsapp.strip(),
                    whatsapp_normalized=normalized_whatsapp,
                    neighborhood=neighborhood.strip(),
                    marketing_consent=marketing_consent,
                    marketing_consent_at=now if marketing_consent else None,
                    marketing_consent_version=(
                        MARKETING_CONSENT_VERSION if marketing_consent else None
                    ),
                )
                db.session.add(current_lead)
                db.session.flush()
            else:
                current_lead.name = first_name
                current_lead.whatsapp = whatsapp.strip()
                current_lead.neighborhood = neighborhood.strip()
                current_lead.marketing_consent = marketing_consent
                current_lead.marketing_consent_at = now if marketing_consent else None
                current_lead.marketing_consent_version = (
                    MARKETING_CONSENT_VERSION if marketing_consent else None
                )
                current_lead.updated_at = now

            preferences = build_declared_preferences(answers)
            registration = Registration(
                lead_id=current_lead.id,
                name=first_name,
                email=None,
                whatsapp=whatsapp.strip(),
                whatsapp_normalized=normalized_whatsapp,
                neighborhood=neighborhood.strip(),
                marketing_consent=marketing_consent,
                marketing_consent_at=now if marketing_consent else None,
                marketing_consent_version=(
                    MARKETING_CONSENT_VERSION if marketing_consent else None
                ),
                marketing_consent_decided_at=now,
                marketing_consent_text_version=MARKETING_CONSENT_VERSION,
                privacy_notice_accepted_at=now,
                privacy_notice_version=PRIVACY_NOTICE_VERSION,
                answers_json=json.dumps(answers, ensure_ascii=False),
                answer_snapshot_json=json.dumps(
                    build_answer_snapshot(answers), ensure_ascii=False
                ),
                score_json=json.dumps(build_score_snapshot(answers), ensure_ascii=False),
                frequency=preferences["frequency"],
                occasion=preferences["occasion"],
                product_interest=preferences["product_interest"],
                experience_preference=preferences["experience_preference"],
                communication_interest=preferences["communication_interest"],
                campaign_source=CAMPAIGN_SOURCE,
                quiz_version=QUIZ_VERSION,
                submission_key=submission_key,
                profile=profile_id,
                created_at=now,
            )
            db.session.add(registration)
            db.session.flush()
            db.session.add(
                MarketingConsentEvent(
                    lead_id=current_lead.id,
                    registration_id=registration.id,
                    consented=marketing_consent,
                    text_version=MARKETING_CONSENT_VERSION,
                    created_at=now,
                )
            )
            db.session.add(SheetSyncJob(registration_id=registration.id))
            db.session.commit()
            return registration

        try:
            registration = save_registration()
        except IntegrityError:
            db.session.rollback()
            existing_registration = Registration.query.filter_by(
                submission_key=submission_key
            ).one_or_none()
            if existing_registration:
                return registration_response(existing_registration)
            existing_lead = Lead.query.filter_by(
                whatsapp_normalized=normalized_whatsapp
            ).one_or_none()
            if existing_lead is None:
                app.logger.exception("Falha de integridade ao salvar lead da Linha Fit.")
                return jsonify(
                    error="Não foi possível salvar seu cadastro. Tente novamente."
                ), 500
            try:
                registration = save_registration(existing_lead)
            except SQLAlchemyError:
                db.session.rollback()
                app.logger.exception("Falha ao salvar cadastro repetido da Linha Fit.")
                return jsonify(
                    error="Não foi possível salvar seu cadastro. Tente novamente."
                ), 500
        except SQLAlchemyError:
            db.session.rollback()
            app.logger.exception("Falha ao salvar cadastro da Linha Fit.")
            return jsonify(
                error="Não foi possível salvar seu cadastro. Tente novamente."
            ), 500

        return registration_response(registration)

    @app.get("/admin")
    @require_admin
    def admin():
        retry_sheet_sync(limit=25)
        leads = Lead.query.order_by(Lead.updated_at.desc()).limit(200).all()
        lead_rows = []
        for lead in leads:
            history = sorted(
                lead.registrations,
                key=lambda registration: registration.created_at,
                reverse=True,
            )
            latest = history[0] if history else None
            lead_rows.append(
                {
                    "lead": lead,
                    "latest": latest,
                    "history_count": len(history),
                    "answers": json.loads(latest.answer_snapshot_json or "{}")
                    if latest
                    else {},
                    "scores": json.loads(latest.score_json or "{}") if latest else {},
                    "history": [
                        {
                            "registration": registration,
                            "answers": json.loads(
                                registration.answer_snapshot_json or "{}"
                            ),
                            "scores": json.loads(registration.score_json or "{}"),
                        }
                        for registration in history
                    ],
                    "preference_labels": {
                        "frequency": option_label("frequency", latest.frequency),
                        "occasion": option_label("occasion", latest.occasion),
                        "product": option_label("product", latest.product_interest),
                        "experience": option_label(
                            "experience", latest.experience_preference
                        ),
                        "news": option_label("news", latest.communication_interest),
                    }
                    if latest
                    else {},
                    "consent_events": sorted(
                        lead.consent_events,
                        key=lambda event: event.created_at,
                        reverse=True,
                    ),
                }
            )
        pending_count = SheetSyncJob.query.filter_by(synced_at=None).count()
        pending_jobs = (
            SheetSyncJob.query.filter_by(synced_at=None)
            .order_by(SheetSyncJob.registration_id.desc())
            .limit(10)
            .all()
        )
        synced_count = SheetSyncJob.query.filter(SheetSyncJob.synced_at.isnot(None)).count()
        return render_template(
            "admin.html",
            lead_rows=lead_rows,
            profiles=ALL_PROFILES,
            sheet_sync_configured=is_configured(app.config),
            pending_jobs=pending_jobs,
            pending_count=pending_count,
            synced_count=synced_count,
        )

    @app.get("/admin/export.xlsx")
    @require_admin
    def export_excel():
        workbook = Workbook()
        contacts_sheet = workbook.active
        contacts_sheet.title = "Contatos"
        answers_sheet = workbook.create_sheet("Respostas")
        summary_sheet = workbook.create_sheet("Resumo")
        header_fill = PatternFill("solid", fgColor="82371F")
        header_font = Font(color="FFFFFF", bold=True)

        def safe_text(value):
            text = str(value or "")
            return f"'{text}" if text.startswith(("=", "+", "-", "@", "\t", "\r")) else text

        def utc_datetime(value):
            if value is None:
                return None
            if value.tzinfo is None:
                value = value.replace(tzinfo=timezone.utc)
            return value.astimezone(timezone.utc).replace(tzinfo=None)

        def add_table(sheet, headers, rows, table_name):
            sheet.append(headers)
            for row in rows:
                sheet.append(row)
            for cell in sheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            if sheet.max_row > 1:
                table = Table(displayName=table_name, ref=sheet.dimensions)
                table.tableStyleInfo = TableStyleInfo(
                    name="TableStyleMedium2",
                    showFirstColumn=False,
                    showLastColumn=False,
                    showRowStripes=True,
                    showColumnStripes=False,
                )
                sheet.add_table(table)
            for column in sheet.columns:
                letter = column[0].column_letter
                width = max(
                    (len(str(cell.value or "")) for cell in column),
                    default=10,
                )
                sheet.column_dimensions[letter].width = min(max(width + 2, 14), 42)

        leads = Lead.query.order_by(Lead.name.asc(), Lead.id.asc()).all()
        registrations = Registration.query.order_by(Registration.created_at.asc()).all()
        registrations_by_lead = {}
        for registration in registrations:
            if registration.lead_id is not None:
                registrations_by_lead.setdefault(registration.lead_id, []).append(registration)

        profile_titles = {
            profile_id: profile["title"] for profile_id, profile in ALL_PROFILES.items()
        }
        contacts_rows = []
        for lead in leads:
            history = registrations_by_lead.get(lead.id, [])
            latest = history[-1] if history else None
            contacts_rows.append(
                [
                    safe_text(lead.name),
                    safe_text(lead.whatsapp),
                    safe_text(lead.neighborhood),
                    profile_titles.get(latest.profile, "—") if latest else "—",
                    len(history),
                    "Sim" if lead.marketing_consent else "Não",
                    utc_datetime(lead.updated_at),
                ]
            )
        add_table(
            contacts_sheet,
            [
                "Nome",
                "WhatsApp",
                "Bairro / cidade",
                "Perfil mais recente",
                "Participações",
                "Aceita novidades no WhatsApp?",
                "Última atualização (UTC)",
            ],
            contacts_rows,
            "Contatos",
        )
        for cell in contacts_sheet["G"][1:]:
            cell.number_format = "dd/mm/yyyy hh:mm"

        answer_rows = []
        for registration in registrations:
            answer_snapshot = json.loads(registration.answer_snapshot_json or "{}")
            answers = json.loads(registration.answers_json)
            response_values = []
            for question in QUESTIONS:
                saved_answer = answer_snapshot.get(question["id"], {}).get("answer")
                if saved_answer is None:
                    saved_answer = option_label(question["id"], answers.get(question["id"]))
                response_values.append(safe_text(saved_answer))
            answer_rows.append(
                [
                    utc_datetime(registration.created_at),
                    safe_text(registration.name),
                    safe_text(registration.whatsapp),
                    safe_text(registration.neighborhood),
                    profile_titles.get(registration.profile, registration.profile),
                    *response_values,
                    "Sim" if registration.marketing_consent else "Não",
                    utc_datetime(registration.marketing_consent_decided_at),
                ]
            )
        add_table(
            answers_sheet,
            [
                "Data e hora (UTC)",
                "Nome",
                "WhatsApp",
                "Bairro / cidade",
                "Perfil",
                *[question["text"] for question in QUESTIONS],
                "Aceita novidades no WhatsApp?",
                "Data da decisão sobre mensagens (UTC)",
            ],
            answer_rows,
            "Respostas",
        )
        for column_name in ("A", "L"):
            for cell in answers_sheet[column_name][1:]:
                cell.number_format = "dd/mm/yyyy hh:mm"

        profile_counts = Counter(
            profile_titles.get(registration.profile, registration.profile)
            for registration in registrations
        )
        opted_in = sum(1 for lead in leads if lead.marketing_consent)
        summary_rows = [
            ["Contatos únicos", len(leads)],
            ["Participações registradas", len(registrations)],
            ["Aceitaram receber novidades", opted_in],
            ["Não aceitaram receber novidades", len(leads) - opted_in],
            ["", ""],
            ["Perfil", "Participações"],
        ]
        summary_rows.extend(
            [
                profile["title"],
                profile_counts.get(profile["title"], 0),
            ]
            for profile in PROFILES.values()
        )
        for row in summary_rows:
            summary_sheet.append(row)
        summary_sheet["A1"].font = Font(size=16, bold=True, color="552616")
        for cell in summary_sheet[6]:
            cell.fill = header_fill
            cell.font = header_font
        summary_sheet.column_dimensions["A"].width = 38
        summary_sheet.column_dimensions["B"].width = 22
        summary_sheet.freeze_panes = "A7"

        output = io.BytesIO()
        workbook.save(output)
        output.seek(0)
        return send_file(
            output,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name="sulamita-linha-fit-relatorio.xlsx",
            max_age=0,
        )

    return app


app = create_app()
