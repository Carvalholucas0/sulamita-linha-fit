import base64
import io
import unittest
from unittest.mock import patch

from app import create_app
from models import Lead, MarketingConsentEvent, Registration, SheetSyncJob, db
from openpyxl import load_workbook
from quiz import QUESTIONS
from sheets import SheetSyncError


class SulamitaAppTests(unittest.TestCase):
    def setUp(self):
        self.app = create_app(
            {
                "TESTING": True,
                "SQLALCHEMY_DATABASE_URI": "sqlite:///:memory:",
                "SECRET_KEY": "test-secret",
                "ADMIN_USER": "test-admin",
                "ADMIN_PASSWORD": "test-password",
            }
        )
        self.client = self.app.test_client()

    def tearDown(self):
        with self.app.app_context():
            db.session.remove()
            db.drop_all()

    def valid_payload(self):
        return {
            "name": "Ana",
            "whatsapp": "11999999999",
            "neighborhood": "São Francisco, Niterói",
            "data_consent": True,
            "marketing_consent": False,
            "submission_key": "11111111-1111-4111-8111-111111111111",
            "answers": {
                question["id"]: question["options"][0]["id"] for question in QUESTIONS
            },
        }

    def test_profile_is_not_available_before_required_registration_submission(self):
        response = self.client.post(
            "/api/result",
            json={"answers": self.valid_payload()["answers"]},
        )

        self.assertEqual(response.status_code, 404)
        with self.app.app_context():
            self.assertEqual(Registration.query.count(), 0)
            self.assertEqual(Lead.query.count(), 0)

    def test_submission_saves_lead_preferences_and_consent_history(self):
        response = self.client.post("/api/submit", json=self.valid_payload())

        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json["success"])
        with self.app.app_context():
            registration = Registration.query.one()
            lead = Lead.query.one()
            consent_event = MarketingConsentEvent.query.one()
            self.assertEqual(registration.name, "Ana")
            self.assertEqual(registration.whatsapp_normalized, "+5511999999999")
            self.assertEqual(registration.neighborhood, "São Francisco, Niterói")
            self.assertEqual(registration.frequency, "weekly")
            self.assertEqual(registration.product_interest, "quiches")
            self.assertEqual(registration.profile, "raiz")
            self.assertEqual(registration.campaign_source, "Evento Linha Fit — São Francisco — 31/10/2026")
            self.assertIn("Toda semana", registration.answer_snapshot_json)
            self.assertFalse(registration.marketing_consent)
            self.assertIsNone(registration.marketing_consent_at)
            self.assertIsNone(registration.marketing_consent_version)
            self.assertIsNotNone(registration.marketing_consent_decided_at)
            self.assertIsNotNone(registration.marketing_consent_text_version)
            self.assertEqual(lead.whatsapp_normalized, "+5511999999999")
            self.assertFalse(lead.marketing_consent)
            self.assertFalse(consent_event.consented)
            self.assertIsNotNone(registration.privacy_notice_accepted_at)
            self.assertIsNotNone(registration.privacy_notice_version)
            self.assertEqual(SheetSyncJob.query.one().registration_id, registration.id)

    def test_repeat_contact_reuses_lead_and_preserves_response_and_consent_history(self):
        first_payload = self.valid_payload()
        self.assertEqual(self.client.post("/api/submit", json=first_payload).status_code, 200)

        second_payload = self.valid_payload()
        second_payload.update(
            {
                "name": "Ana Clara",
                "whatsapp": "+55 (11) 99999-9999",
                "neighborhood": "Icaraí, Niterói",
                "marketing_consent": True,
                "submission_key": "22222222-2222-4222-8222-222222222222",
            }
        )
        second_payload["answers"]["product"] = "sweets"
        response = self.client.post("/api/submit", json=second_payload)

        self.assertEqual(response.status_code, 200)
        with self.app.app_context():
            self.assertEqual(Lead.query.count(), 1)
            self.assertEqual(Registration.query.count(), 2)
            self.assertEqual(MarketingConsentEvent.query.count(), 2)
            lead = Lead.query.one()
            self.assertEqual(lead.name, "Ana")
            self.assertEqual(lead.neighborhood, "Icaraí, Niterói")
            self.assertTrue(lead.marketing_consent)
            self.assertEqual(
                [event.consented for event in sorted(lead.consent_events, key=lambda event: event.id)],
                [False, True],
            )
            accepted_registration = max(lead.registrations, key=lambda item: item.id)
            declined_registration = min(lead.registrations, key=lambda item: item.id)
            self.assertIsNotNone(accepted_registration.marketing_consent_at)
            self.assertIsNone(declined_registration.marketing_consent_at)
            products = {registration.product_interest for registration in lead.registrations}
            self.assertEqual(products, {"quiches", "sweets"})

        credentials = base64.b64encode(b"test-admin:test-password").decode()
        history_response = self.client.get(
            "/admin",
            headers={"Authorization": f"Basic {credentials}"},
        )
        history_html = history_response.get_data(as_text=True)
        self.assertIn("Cadastro #1", history_html)
        self.assertIn("Cadastro #2", history_html)

    def test_retry_with_same_submission_key_does_not_duplicate_registration(self):
        payload = self.valid_payload()
        first = self.client.post("/api/submit", json=payload)
        retry = self.client.post("/api/submit", json=payload)

        self.assertEqual(first.status_code, 200)
        self.assertEqual(retry.status_code, 200)
        with self.app.app_context():
            self.assertEqual(Lead.query.count(), 1)
            self.assertEqual(Registration.query.count(), 1)
            self.assertEqual(MarketingConsentEvent.query.count(), 1)

    def test_submission_is_saved_when_sheet_sync_fails_and_admin_retries(self):
        self.app.config.update(
            SHEETS_SERVICE_ACCOUNT_JSON="{}",
            SHEETS_SPREADSHEET_ID="test-sheet",
        )
        with patch("app.sync_registration", side_effect=SheetSyncError("temporary outage")):
            response = self.client.post("/api/submit", json=self.valid_payload())

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json["sheet_backup"], "pending")
        with self.app.app_context():
            registration = Registration.query.one()
            job = SheetSyncJob.query.one()
            self.assertEqual(registration.name, "Ana")
            self.assertIsNone(job.synced_at)
            self.assertEqual(job.attempt_count, 1)

        credentials = base64.b64encode(b"test-admin:test-password").decode()
        with patch("app.sync_registration") as sync:
            response = self.client.get(
                "/admin",
                headers={"Authorization": f"Basic {credentials}"},
            )

        self.assertEqual(response.status_code, 200)
        self.assertIn("Planilha sincronizada", response.get_data(as_text=True))
        sync.assert_called_once()
        with self.app.app_context():
            job = SheetSyncJob.query.one()
            self.assertIsNotNone(job.synced_at)
            self.assertEqual(job.attempt_count, 2)

    def test_production_submission_queues_sheet_sync_without_waiting(self):
        self.app.testing = False
        self.app.config.update(
            SHEETS_SERVICE_ACCOUNT_JSON="{}",
            SHEETS_SPREADSHEET_ID="test-sheet",
        )
        with patch("app.ThreadPoolExecutor.submit") as submit:
            response = self.client.post("/api/submit", json=self.valid_payload())

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json["sheet_backup"], "pending")
        submit.assert_called_once()
        with self.app.app_context():
            self.assertEqual(Registration.query.count(), 1)
            self.assertIsNone(SheetSyncJob.query.one().synced_at)

    def test_submission_validates_phone_consent_and_all_answers(self):
        payload = self.valid_payload()
        payload["whatsapp"] = "12345"
        response = self.client.post("/api/submit", json=payload)
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.json["field"], "whatsapp")

        payload = self.valid_payload()
        payload["data_consent"] = False
        response = self.client.post("/api/submit", json=payload)
        self.assertEqual(response.status_code, 400)

        payload = self.valid_payload()
        payload["answers"].pop(QUESTIONS[0]["id"])
        self.assertEqual(self.client.post("/api/submit", json=payload).status_code, 400)
        with self.app.app_context():
            self.assertEqual(Registration.query.count(), 0)
            self.assertEqual(Lead.query.count(), 0)

    def test_admin_and_excel_export_require_authentication_and_summarize_leads(self):
        self.assertEqual(self.client.get("/admin").status_code, 401)
        self.assertEqual(self.client.get("/admin/export.xlsx").status_code, 401)
        first_payload = self.valid_payload()
        first_payload["answers"]["occasion"] = "after_workout"
        self.client.post("/api/submit", json=first_payload)
        second_payload = self.valid_payload()
        second_payload.update(
            {
                "name": "=1+1",
                "marketing_consent": True,
                "submission_key": "22222222-2222-4222-8222-222222222222",
            }
        )
        self.client.post("/api/submit", json=second_payload)

        credentials = base64.b64encode(b"test-admin:test-password").decode()
        headers = {"Authorization": f"Basic {credentials}"}
        admin_response = self.client.get("/admin", headers=headers)
        self.assertEqual(admin_response.status_code, 200)
        self.assertIn("São Francisco, Niterói", admin_response.get_data(as_text=True))

        response = self.client.get("/admin/export.xlsx", headers=headers)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.mimetype,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        workbook = load_workbook(io.BytesIO(response.data), data_only=True)
        self.assertEqual(workbook.sheetnames, ["Contatos", "Respostas", "Resumo"])

        contacts = workbook["Contatos"]
        self.assertEqual(contacts.max_row, 2)
        self.assertEqual(contacts["A2"].value, "'=1+1")
        self.assertEqual(contacts["B2"].value, "11999999999")
        self.assertEqual(contacts["C2"].value, "São Francisco, Niterói")
        self.assertEqual(contacts["E2"].value, 2)
        self.assertEqual(contacts["F2"].value, "Sim")
        self.assertEqual(contacts["A2"].data_type, "s")

        answers = workbook["Respostas"]
        self.assertEqual(answers.max_row, 3)
        self.assertEqual(answers["G2"].value, "Depois da corrida, caminhada ou treino. Afinal, eu treino pra isso!")
        self.assertEqual(answers["K2"].value, "Não")
        self.assertEqual(answers["K3"].value, "Sim")
        self.assertEqual(answers["B3"].value, "'=1+1")

        summary = workbook["Resumo"]
        self.assertEqual(summary["B1"].value, 1)
        self.assertEqual(summary["B2"].value, 2)
        self.assertEqual(summary["B3"].value, 1)
        self.assertEqual(summary["B4"].value, 0)

    def test_landing_and_privacy_pages_render(self):
        response = self.client.get("/")
        self.assertEqual(response.status_code, 200)
        page = response.get_data(as_text=True)
        self.assertIn("Qual SulaLover é você?", page)
        self.assertIn("preencha seus dados e responda às cinco perguntas", page)
        self.assertNotIn("cadastro é opcional", page)
        self.assertEqual(self.client.get("/privacidade").status_code, 200)
